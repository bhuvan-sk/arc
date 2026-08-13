"""DOCX exporter using python-docx."""
import os
from typing import Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_docx(session: dict[str, Any], output_path: str) -> int:
    """Generate Word (.docx) export at output_path. Returns file size in bytes."""
    doc = Document()

    title = session.get("title", "Architecture Explanation")
    subtitle = session.get("subtitle", "")
    layers = session.get("layers", [])
    unlocked = session.get("unlocked_index", 0)

    # Title
    p_title = doc.add_heading(level=0)
    run_title = p_title.add_run(title)
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 15, 17)

    if subtitle:
        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run(subtitle)
        run_sub.font.name = "Arial"
        run_sub.font.size = Pt(11)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(100, 100, 110)

    doc.add_paragraph()  # spacing

    for l in layers:
        idx = l.get("index", 0)
        if idx > unlocked:
            continue

        l_name = l.get("name", f"Layer {idx+1}")
        band_label = l.get("band_label", "")
        narrative = l.get("narrative", {}) or {}
        n_title = narrative.get("title", l_name)
        n_body = narrative.get("body", "")
        items = narrative.get("items", [])
        decisions = l.get("scale_decisions", [])

        # Kicker
        p_kick = doc.add_paragraph()
        r_kick = p_kick.add_run(band_label)
        r_kick.font.name = "Arial"
        r_kick.font.size = Pt(9)
        r_kick.font.bold = True
        r_kick.font.color.rgb = RGBColor(59, 111, 229)

        # Layer title
        h2 = doc.add_heading(level=1)
        r_h2 = h2.add_run(n_title)
        r_h2.font.name = "Arial"
        r_h2.font.size = Pt(16)
        r_h2.font.color.rgb = RGBColor(17, 17, 19)

        # Body
        p_body = doc.add_paragraph(n_body)
        p_body.paragraph_format.space_after = Pt(12)

        # Items
        if items:
            p_comp = doc.add_paragraph()
            r_comp = p_comp.add_run("Components:")
            r_comp.font.bold = True

            for it in items:
                p_it = doc.add_paragraph(style="List Bullet")
                r_it_title = p_it.add_run(it.get("title", ""))
                r_it_title.font.bold = True
                if it.get("sub"):
                    p_it.add_run(f" — {it.get('sub')}")

        # Decisions
        if decisions:
            p_dec_title = doc.add_paragraph()
            r_dec_title = p_dec_title.add_run("Scale Decisions:")
            r_dec_title.font.bold = True

            for d in decisions:
                p_d = doc.add_paragraph(style="List Bullet")
                p_d.add_run(f"[{d.get('concern','').upper()}] ").bold = True
                p_d.add_run(d.get("decision", ""))
                if d.get("tradeoff_seed"):
                    p_d.add_run(f" (Trade-off: {d.get('tradeoff_seed')})").italic = True

        doc.add_paragraph()  # section gap

    doc.save(output_path)
    return os.path.getsize(output_path)
