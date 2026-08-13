"""DOCX ingestion — extract paragraph and table text via python-docx."""
import io
from docx import Document


def extract_docx_text(docx_bytes: bytes) -> str:
    """Extract all text from a Word document."""
    doc = Document(io.BytesIO(docx_bytes))
    lines: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines)
